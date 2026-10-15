"""
PDF tables to an Excel workbook (part of the ZapConvert media worker, AGPL-3.0; see LICENSE).

Every table PyMuPDF detects becomes its own sheet ("P2 T1" = page 2, table 1). Pages without a
detectable table fall back to their text lines, split into columns on runs of 2+ spaces, so
simple column layouts still land in cells. Numbers become real numbers unless a leading zero or
plus sign means they are identifiers (postcodes, phone numbers).

usage: python3 pdf_xlsx.py input.pdf output.xlsx
exit 3 when the PDF has no extractable text (scanned images).
"""

import re
import sys

import pymupdf as fitz
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

NUMBER = re.compile(r"^-?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$")
MAX_SHEETS = 250


def cell(value):
    if value is None:
        return None
    text = " ".join(str(value).split())
    if not text:
        return None
    if NUMBER.match(text) and not re.match(r"^-?0\d", text):
        number = float(text.replace(",", ""))
        return int(number) if number.is_integer() and "." not in text else number
    return text


def add_sheet(book, title, rows):
    sheet = book.create_sheet(title=title[:31])
    widths = {}
    for row in rows:
        values = [cell(v) for v in row]
        sheet.append(values)
        for i, v in enumerate(values, 1):
            if v is not None:
                widths[i] = max(widths.get(i, 0), min(60, len(str(v))))
    for i, w in widths.items():
        sheet.column_dimensions[get_column_letter(i)].width = max(8, w + 2)


def main(src, dst):
    doc = fitz.open(src)
    if doc.needs_pass:
        print("password-protected", file=sys.stderr)
        sys.exit(4)
    book = Workbook()
    book.remove(book.active)
    any_text = False
    for page in doc:
        if len(book.sheetnames) >= MAX_SHEETS:
            break
        text = page.get_text("text")
        if text.strip():
            any_text = True
        tables = page.find_tables().tables
        if tables:
            for t, table in enumerate(tables, 1):
                rows = [r for r in table.extract() if any(c not in (None, "") for c in r)]
                if rows:
                    add_sheet(book, f"P{page.number + 1} T{t}", rows)
        elif text.strip():
            rows = [re.split(r"\s{2,}", line.strip()) for line in text.splitlines() if line.strip()]
            add_sheet(book, f"P{page.number + 1} Text", rows)
    if not any_text:
        print("no text layer", file=sys.stderr)
        sys.exit(3)
    book.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
